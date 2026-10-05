from datetime import date, datetime

import openpyxl

from procurement_search.attribute_extractor import Quantity
from procurement_search.export import export_companies_to_excel
from procurement_search.models import (
    ACTUAL_ADDRESS,
    Availability,
    AvailabilityStatus,
    Company,
    Evidence,
    FieldValue,
    IntentType,
    LEGAL_ADDRESS,
    PageType,
    ProductMatch,
    ProductOffer,
    VerificationFlag,
)
from procurement_search.quantity_match import Verdict


def _company(name: str, with_availability: bool = False) -> Company:
    company = Company(
        inn=None, ogrn=None, name=FieldValue(name, "test", date(2026, 1, 1)), status="неизвестно"
    )
    if with_availability:
        company.availability = Availability(
            status=AvailabilityStatus.IN_STOCK_QTY,
            quantity=Quantity(10.0, "шт", "10 шт"),
            pack_size=None,
            min_order=None,
            lead_time_days=None,
            price=None,
            source_url="https://x.example/product",
            checked_at=datetime(2026, 8, 6, 12, 0),
            evidence="в наличии: 10 шт",
        )
        company.availability_verdict = Verdict.NOT_ENOUGH.value
        company.availability_verdict_text = "Недостаточно — есть всего 10 шт из 11"
    return company


def test_export_without_availability_keeps_header_at_row_one(tmp_path):
    output = export_companies_to_excel([_company("ООО Тест")], tmp_path / "out.xlsx")

    wb = openpyxl.load_workbook(output)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Компания"
    assert ws.cell(row=2, column=1).value == "ООО Тест"


def test_export_service_intent_uses_service_headers(tmp_path):
    output = export_companies_to_excel(
        [_company("ООО Утилизация")],
        tmp_path / "service.xlsx",
        intent_type=IntentType.SERVICE,
    )

    wb = openpyxl.load_workbook(output)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    assert headers[0] == "Исполнитель"
    assert "Соответствие услуге" in headers
    assert "Наличие товара" not in headers
    assert "Цена / тариф" in headers


def test_export_product_intent_uses_offer_price_when_company_price_missing(tmp_path):
    company = _company("ООО Оффер")
    company.product_offers.append(
        ProductOffer(
            company_id=None,
            landing_url="https://shop.example/product",
            page_type=PageType.PRODUCT_DETAIL,
            product_match=ProductMatch.COMPATIBLE,
            price=50000.0,
            currency="RUB",
            evidence=[
                Evidence(
                    field="price",
                    value="50 000 руб.",
                    source_url="https://shop.example/product",
                    source_text="Цена 50 000 руб.",
                )
            ],
        )
    )

    output = export_companies_to_excel(
        [company],
        tmp_path / "offer-price.xlsx",
        intent_type=IntentType.PRODUCT,
    )

    wb = openpyxl.load_workbook(output)
    ws = wb.active
    assert ws.cell(row=2, column=23).value == "50 000 руб."
    assert "shop.example/product" in ws.cell(row=2, column=24).value


def test_export_with_summary_shifts_header_down_and_writes_summary_row(tmp_path):
    output = export_companies_to_excel(
        [_company("ООО Тест")],
        tmp_path / "out.xlsx",
        required_qty=Quantity(11.0, "шт", "11 шт"),
        summary="Требуется: 11 шт\nПодтверждено на складах: 10 шт у 1 поставщиков",
    )

    wb = openpyxl.load_workbook(output)
    ws = wb.active
    assert "Требуется: 11 шт" in ws.cell(row=1, column=1).value
    assert ws.cell(row=2, column=1).value == "Компания"
    assert ws.cell(row=3, column=1).value == "ООО Тест"


def test_export_writes_availability_columns():
    """Индексы новых колонок (15-22, см. export._COL_*): "Запрошено",
    "Найдено", ..., "Наличие (кол-во): источник/дата"."""
    from procurement_search.export import (
        _COL_AVAILABILITY_EVIDENCE,
        _COL_AVAILABILITY_SOURCE,
        _COL_AVAILABILITY_VERDICT,
        _COL_FOUND_QTY,
        _COL_REQUESTED,
    )

    import tempfile
    from pathlib import Path

    with tempfile.TemporaryDirectory() as tmp:
        output = export_companies_to_excel(
            [_company("ООО Тест", with_availability=True)],
            Path(tmp) / "out.xlsx",
            required_qty=Quantity(11.0, "шт", "11 шт"),
        )
        wb = openpyxl.load_workbook(output)
        ws = wb.active

        assert ws.cell(row=2, column=_COL_REQUESTED).value == "11 шт"
        assert ws.cell(row=2, column=_COL_FOUND_QTY).value == "10 шт"
        assert ws.cell(row=2, column=_COL_AVAILABILITY_VERDICT).value == (
            "Недостаточно — есть всего 10 шт из 11"
        )
        assert ws.cell(row=2, column=_COL_AVAILABILITY_EVIDENCE).value == "в наличии: 10 шт"
        assert "x.example" in ws.cell(row=2, column=_COL_AVAILABILITY_SOURCE).value


def test_export_without_availability_leaves_new_columns_blank(tmp_path):
    from procurement_search.export import _COL_AVAILABILITY_VERDICT, _COL_FOUND_QTY

    output = export_companies_to_excel([_company("ООО Тест")], tmp_path / "out.xlsx")

    wb = openpyxl.load_workbook(output)
    ws = wb.active
    # openpyxl не различает "" и пустую ячейку при перечитывании файла —
    # обе читаются как None, хотя записывались пустой строкой (см.
    # export.py: "" для отсутствующих полей, тот же принцип, что и у
    # остальных пустых колонок в этом экспортере).
    assert ws.cell(row=2, column=_COL_FOUND_QTY).value is None
    assert ws.cell(row=2, column=_COL_AVAILABILITY_VERDICT).value is None


def test_export_shows_actual_address_over_legal_egrul_address(tmp_path):
    """Превью адреса показывает ФАКТИЧЕСКИЙ адрес работы (Казань), а не
    юридический адрес ЕГРЮЛ (Барнаул) — у kazan.geogrunt.ru головной офис в
    Барнауле, но компания работает в Казани (design-обсуждение: юридический
    адрес закладывается в карточку как доп. источник, но не подменяет место
    работы на превью)."""
    company = _company("ООО Геогрунт")
    company.contacts["address"] = [
        FieldValue(
            "г Казань, ул Сибирский Тракт, д 39, помещ 1002",
            "текст сайта (Слой 2)",
            date(2026, 8, 7),
            VerificationFlag.UNVERIFIED,
            kind=ACTUAL_ADDRESS,
        ),
        FieldValue(
            "656031, АЛТАЙСКИЙ КРАЙ, Г.О. ГОРОД БАРНАУЛ, Г. БАРНАУЛ, УЛ. ПРИВОКЗАЛЬНАЯ, Д. 49",
            "ЕГРЮЛ (Dadata)",
            date(2026, 8, 7),
            VerificationFlag.CONFIRMED,
            kind=LEGAL_ADDRESS,
        ),
    ]

    output = export_companies_to_excel([company], tmp_path / "out.xlsx")
    wb = openpyxl.load_workbook(output)
    ws = wb.active
    # Адрес — колонка 8, источник/дата адреса — колонка 9
    assert "Казань" in ws.cell(row=2, column=8).value
    assert "Барнаул" not in ws.cell(row=2, column=8).value

    # Наоборот, когда фактического адреса нет — показывается юридический из ЕГРЮЛ
    company2 = _company("ООО Только юр.адрес")
    company2.contacts["address"] = [
        FieldValue(
            "г. Москва, ул. Ленина, 1",
            "ЕГРЮЛ (Dadata)",
            date(2026, 8, 7),
            VerificationFlag.CONFIRMED,
            kind=LEGAL_ADDRESS,
        ),
    ]
    output2 = export_companies_to_excel([company2], tmp_path / "out2.xlsx")
    wb2 = openpyxl.load_workbook(output2)
    ws2 = wb2.active
    assert "Москва" in ws2.cell(row=2, column=8).value
