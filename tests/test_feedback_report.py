"""Тесты feedback_report.py и фильтрации по датам в feedback.FeedbackStore (без сети)."""

from datetime import datetime, timezone

from openpyxl import load_workbook

from procurement_search import feedback, feedback_report


def _make_store(tmp_path):
    db_path = tmp_path / "feedback_test.db"
    store = feedback.FeedbackStore(db_path)
    return db_path, store


def test_list_all_date_filter(tmp_path):
    """list_all(since/until) фильтрует по ISO-датам; until включает весь день."""
    db, store = _make_store(tmp_path)
    # Кладём записи напрямую с известным created_at, чтобы проверить фильтр.
    samples = {26: "up", 28: "down", 30: "text"}
    for idx, (day, rating) in enumerate(samples.items(), start=1):
        store._conn.execute(
            "INSERT INTO feedback (report_key, query, name, email, rating, comment, created_at) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (f"r{idx}", "запрос", f"Имя{idx}", f"u{idx}@x.ru", rating, "комм",
             f"2026-09-{day:02d}T08:00:00+00:00"),
        )
    store._conn.commit()

    week = store.list_all(since="2026-09-21", until="2026-09-27")
    assert [r["report_key"] for r in week] == ["r1"]  # только 26-е

    whole_day = store.list_all(since="2026-09-28", until="2026-09-30")
    # до конца 30-го включительно
    assert {r["report_key"] for r in whole_day} == {"r2", "r3"}


def test_group_by_week_sorts_and_buckets():
    rows = [
        {"created_at": "2026-09-28T08:00:00+00:00", "rating": "up", "comment": ""},
        {"created_at": "2026-09-30T08:00:00+00:00", "rating": "down", "comment": "плохо"},
        {"created_at": "2026-09-21T08:00:00+00:00", "rating": "up", "comment": ""},
        {"created_at": "bad-date", "rating": "up", "comment": ""},
    ]
    weekly = feedback_report.group_by_week(rows)
    # bad-date пропускается; недели 28.09 и 21.09, сортировка по убыванию
    assert len(weekly) == 2
    assert weekly[0][0].day == 28
    assert weekly[1][0].day == 21
    assert len(weekly[0][1]) == 2  # 28 и 30 — одна неделя


def test_week_start_monday():
    w = feedback_report._week_start(datetime(2026, 9, 30, 8, 0, tzinfo=timezone.utc))
    assert w.isoweekday() == 1


def test_export_creates_two_sheets(tmp_path):
    rows = [
        {"created_at": "2026-09-28T08:00:00+00:00", "query": "q", "name": "Иван",
         "email": "i@x.ru", "rating": "up", "comment": ""},
        {"created_at": "2026-09-25T08:00:00+00:00", "query": "q", "name": "Петя",
         "email": "p@x.ru", "rating": "text", "comment": "добавьте счёт"},
    ]
    out = feedback_report.export_feedback_to_excel(rows, tmp_path / "out.xlsx")
    assert out.is_file()
    wb = load_workbook(out)
    assert wb.sheetnames == ["Все отзывы", "Сводка по неделям"]
    all_ws = wb["Все отзывы"]
    assert all_ws.max_row == 3  # шапка + 2 данных
    summary = wb["Сводка по неделям"]
    assert summary.max_row >= 3  # шапка + недели + ИТОГО


def test_summary_counts_and_total(tmp_path):
    """Сводка корректно считает +/-/комментарии и строку ИТОГО."""
    rows = [
        {"created_at": "2026-09-28T08:00:00+00:00", "query": "q", "name": "A",
         "email": "a@x.ru", "rating": "up", "comment": "хорошо"},
        {"created_at": "2026-09-29T08:00:00+00:00", "query": "q", "name": "B",
         "email": "b@x.ru", "rating": "down", "comment": ""},
        {"created_at": "2026-09-22T08:00:00+00:00", "query": "q", "name": "C",
         "email": "c@x.ru", "rating": "text", "comment": "комм"},
    ]
    out = feedback_report.export_feedback_to_excel(rows, tmp_path / "s.xlsx")
    wb = load_workbook(out)
    s = wb["Сводка по неделям"]
    last = s.max_row
    total = [s.cell(row=last, column=c).value for c in range(1, 9)]
    assert total[0] == "ИТОГО"
    assert total[3] == 3  # всего
    assert total[4] == 1  # 👍
    assert total[5] == 1  # 👎
    assert total[6] == 1  # комментарии (rating=text)
    assert total[7] == 2  # отзывов с текстом (A и C)


def test_build_feedback_report_end_to_end(tmp_path):
    db, store = _make_store(tmp_path)
    store.add("r1", "запрос", "Иван", "i@x.ru", "up", "супер")
    out = feedback_report.build_feedback_report(tmp_path / "report.xlsx", db_path=db)
    assert out.is_file()
    wb = load_workbook(out)
    assert "Все отзывы" in wb.sheetnames
    assert "Сводка по неделям" in wb.sheetnames


def test_send_feedback_report_email_needs_smtp(monkeypatch, tmp_path):
    """Без настроенной почты отправка молча возвращает False, не падает."""
    monkeypatch.setattr(feedback, "load_smtp_config", lambda: None)
    out = feedback_report.export_feedback_to_excel([], tmp_path / "e.xlsx")
    assert feedback_report.send_feedback_report_email(out) is False