"""Шаг [7] пайплайна: экспорт в Excel (design_doc §8).

Принцип "фифти-фифти" реализован буквально: для каждого контактного поля
экспортируется не только значение, но источник, дата получения и флаг
достоверности — так байер видит, какие именно ячейки требуют проверки,
вместо перепроверки всей строки.
"""

from __future__ import annotations

import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from procurement_search.attribute_extractor import Quantity
from procurement_search.models import Company, StockStatus, VerificationFlag
from procurement_search.quantity_match import Verdict

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

_FLAG_FILL = {
    VerificationFlag.CONFIRMED: PatternFill(
        start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"
    ),
    VerificationFlag.UNVERIFIED: PatternFill(
        start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"
    ),
    VerificationFlag.STALE: PatternFill(
        start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"
    ),
}

# Отдельная палитра от _FLAG_FILL: NOT_CHECKED — это не "недостоверно", а
# "не проверялось вовсе" (Слой 3 не запускался — deep_relevance/
# relevance_llm_check выключены или LLM недоступна), поэтому нейтральный
# серый, а не жёлтый "не проверен" из VerificationFlag.
_STOCK_FILL = {
    StockStatus.IN_STOCK: PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
    StockStatus.CLARIFY: PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    StockStatus.OUT_OF_STOCK: PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    StockStatus.NOT_CHECKED: PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid"),
}

# Availability (см. models.py) — отдельная от _STOCK_FILL палитра, по
# quantity_match.Verdict, а не models.StockStatus (см. их докстринги про
# мотивацию держать обе проверки параллельно). Зелёный — вопрос закрыт,
# жёлтый — нужен звонок/уточнение, серый — данных нет, красный — сюда
# вообще не должно доходить (OUT_OF_STOCK — knockout в pipeline.py, но
# заливка на всякий случай определена, а не оставлена без стиля).
_AVAILABILITY_FILL = {
    Verdict.ENOUGH.value: PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
    Verdict.NOT_ENOUGH.value: PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    Verdict.ON_ORDER.value: PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    Verdict.IN_STOCK_NO_QTY.value: PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid"),
    Verdict.UNIT_MISMATCH.value: PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid"),
    Verdict.UNKNOWN.value: PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid"),
    Verdict.OUT_OF_STOCK.value: PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
}

_COLUMNS = [
    ("Компания", 32),
    ("ИНН", 14),
    ("Статус", 18),
    ("Телефон", 20),
    ("Телефон: источник/дата", 26),
    ("Email", 24),
    ("Email: источник/дата", 26),
    ("Адрес", 30),
    ("Адрес: источник/дата", 26),
    ("Сайт", 26),
    ("Сайт: источник/дата", 26),
    ("Источники (каталоги)", 20),
    ("Наличие товара", 16),
    ("Наличие товара: цитата с сайта", 30),
    ("Запрошено", 14),
    ("Найдено", 16),  # штук ИЛИ метров — см. pipeline.SearchResult.effective_order_amount
    ("Фасовка", 16),
    ("Мин. партия", 14),
    ("Срок поставки", 14),
    ("Сравнение с заказом", 30),
    ("Наличие (кол-во): подтверждение", 30),
    ("Наличие (кол-во): источник/дата", 26),
    ("Цена", 16),
    ("Цена: источник/дата", 26),
    ("Score: релевантность", 12),
    ("Score: доверие", 12),
    ("Score: полнота данных", 14),
    ("Score: итого", 12),
]

# Индексы (1-based) новых колонок — см. _COLUMNS выше, "Запрошено" первая
# из восьми, вставленных между "Наличие товара: цитата с сайта" (14) и
# "Цена" (сдвинулась с 15 на 23).
_COL_REQUESTED = 15
_COL_FOUND_QTY = 16
_COL_PACK_SIZE = 17
_COL_MIN_ORDER = 18
_COL_LEAD_TIME = 19
_COL_AVAILABILITY_VERDICT = 20
_COL_AVAILABILITY_EVIDENCE = 21
_COL_AVAILABILITY_SOURCE = 22
_COL_PRICE = 23
_COL_PRICE_SOURCE = 24


def _format_quantity(q: Quantity | None) -> str:
    return f"{q.value:g} {q.unit}" if q is not None else ""


def _field_summary(field_values) -> tuple[str, str, VerificationFlag]:
    if not field_values:
        return "", "", VerificationFlag.UNVERIFIED
    fv = field_values[0]
    return fv.value, f"{fv.source}, {fv.retrieved_at.isoformat()}", fv.confidence


def _single_field_summary(fv) -> tuple[str, str, VerificationFlag]:
    """Как _field_summary, но для одиночного поля (не списка) — см.
    Company.price в models.py: у цены, в отличие от contacts, нет
    нескольких источников на одного кандидата."""
    if fv is None:
        return "", "", VerificationFlag.UNVERIFIED
    return fv.value, f"{fv.source}, {fv.retrieved_at.isoformat()}", fv.confidence


def _write_header(ws: Worksheet, header_row: int) -> None:
    for col_idx, (title, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width


def export_companies_to_excel(
    companies: list[Company],
    output_path: str | Path,
    required_qty: Quantity | None = None,
    summary: str | None = None,
) -> Path:
    """`required_qty`/`summary` — опциональны (см. pipeline.summarize_availability):
    без них поведение экспорта не меняется (кроме сдвига колонок 15-20 на
    23-28, см. _COLUMNS) — те 8 новых колонок наличия просто пустые, если
    check_availability не использовался. `summary` — одна строка сводки
    над шапкой ("Требуется: 11 шт, подтверждено на складах: ..."), сдвигает
    шапку и данные на одну строку вниз."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Поставщики"

    header_row = 1
    if summary:
        summary_cell = ws.cell(row=1, column=1, value=summary)
        summary_cell.font = Font(bold=True)
        summary_cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_COLUMNS))
        header_row = 2

    _write_header(ws, header_row)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate

    # Порядок берётся как есть, БЕЗ пересортировки по чистому score.total —
    # вызывающий код (pipeline.search_and_score) уже отсортировал компании
    # правильно, включая нечисловые правила ранжирования поверх score
    # (например, деприоритизация маркетплейсов в pipeline._ranking_key,
    # которая должна пересиливать более высокий score — пересортировка
    # здесь тихо стирала бы такие правила, реально это уже случалось).
    for row_idx, company in enumerate(companies, start=header_row + 1):
        phone_val, phone_src, phone_flag = _field_summary(company.contacts.get("phone", []))
        email_val, email_src, email_flag = _field_summary(company.contacts.get("email", []))
        addr_val, addr_src, addr_flag = _field_summary(company.contacts.get("address", []))
        site_val, site_src, site_flag = _field_summary(company.contacts.get("website", []))
        price_val, price_src, price_flag = _single_field_summary(company.price)

        availability = company.availability
        found_val = _format_quantity(availability.quantity) if availability else ""
        pack_val = _format_quantity(availability.pack_size) if availability else ""
        min_order_val = _format_quantity(availability.min_order) if availability else ""
        lead_time_val = (
            availability.lead_time_days if availability and availability.lead_time_days is not None else ""
        )
        availability_src_val = (
            f"{availability.source_url}, {availability.checked_at.date().isoformat()}"
            if availability
            else ""
        )

        score = company.score
        row = [
            company.name.value,
            company.inn or "",
            company.status,
            phone_val,
            phone_src,
            email_val,
            email_src,
            addr_val,
            addr_src,
            site_val,
            site_src,
            ", ".join(company.sources),
            company.stock_status.value,
            company.stock_status_quote or "",
            _format_quantity(required_qty),
            found_val,
            pack_val,
            min_order_val,
            lead_time_val,
            company.availability_verdict_text or "",
            availability.evidence if availability and availability.evidence else "",
            availability_src_val,
            price_val,
            price_src,
            round(score.relevance, 3) if score else "",
            round(score.trust, 3) if score else "",
            round(score.confidence, 3) if score else "",
            round(score.total, 3) if score else "",
        ]
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

        for col_idx, flag in (
            (4, phone_flag),
            (6, email_flag),
            (8, addr_flag),
            (10, site_flag),
            (_COL_PRICE, price_flag),
        ):
            ws.cell(row=row_idx, column=col_idx).fill = _FLAG_FILL[flag]

        # "Наличие товара" — 13-я колонка (см. _COLUMNS): после "Источники (каталоги)".
        ws.cell(row=row_idx, column=13).fill = _STOCK_FILL[company.stock_status]

        if company.availability_verdict is not None:
            ws.cell(row=row_idx, column=_COL_AVAILABILITY_VERDICT).fill = _AVAILABILITY_FILL[
                company.availability_verdict
            ]

    output_path = Path(output_path)
    # Сохранить во временный файл и os.replace, а не wb.save(output_path)
    # напрямую — та же атомарность, что и у webapp._save_report_index_unlocked:
    # если процесс убьют посреди записи (OOM/деплой), на диске останется
    # либо старой версии не было вовсе, либо целый новый файл, но не
    # обрубленный битый .xlsx под именем, которое уже попало в историю.
    tmp_path = output_path.with_name(output_path.name + ".tmp")
    wb.save(tmp_path)
    os.replace(tmp_path, output_path)
    return output_path
